from __future__ import annotations

import logging
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

import requests
from django.conf import settings
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError

from apps.accounts.models import Employee
from apps.assignments.models import Assignment
from apps.attendance.models import Attendance, Session
from apps.common.cloudinary_service import upload_attendance_image, upload_profile_image
from apps.common.utils import distance_meters
from apps.tracking.models import LocationLog
from apps.vision.services import FaceService, LivenessService

logger = logging.getLogger(__name__)

face_service = FaceService()
liveness_service = LivenessService()


def get_active_assignment(employee: Employee) -> Assignment | None:
    today = timezone.localdate()
    return (
        Assignment.objects.filter(employee=employee, visit_date=today)
        .exclude(status=Assignment.Status.CANCELLED)
        .order_by("-created_at")
        .first()
    )


def geocode_address(address: str) -> tuple[float, float] | None:
    if not address:
        return None
    try:
        response = requests.get(
            "https://nominatim.openstreetmap.org/search",
            params={"q": address, "format": "jsonv2", "limit": 1},
            headers={"User-Agent": "EmployeeHub/1.0"},
            timeout=10,
        )
        response.raise_for_status()
        data = response.json()
        if not data:
            return None
        return float(data[0]["lat"]), float(data[0]["lon"])
    except Exception:
        return None


def ensure_default_coordinates(employee: Employee) -> bool:
    if employee.default_latitude is not None and employee.default_longitude is not None:
        return True
    if not employee.default_address:
        return False
    coords = geocode_address(employee.default_address)
    if not coords:
        return False
    employee.default_latitude, employee.default_longitude = coords
    employee.save(update_fields=["default_latitude", "default_longitude"])
    return True


def reverse_geocode_coordinates(latitude: float, longitude: float) -> str | None:
    if latitude is None or longitude is None:
        return None
    try:
        response = requests.get(
            "https://nominatim.openstreetmap.org/reverse",
            params={"lat": latitude, "lon": longitude, "format": "jsonv2"},
            headers={"User-Agent": "EmployeeHub/1.0"},
            timeout=10,
        )
        response.raise_for_status()
        data = response.json()
        if data and "display_name" in data:
            return data["display_name"]
        return None
    except Exception:
        return None


def ensure_default_address(employee: Employee) -> bool:
    if employee.default_address and employee.default_address.strip():
        return True
    if employee.default_latitude is None or employee.default_longitude is None:
        return False
    address = reverse_geocode_coordinates(float(employee.default_latitude), float(employee.default_longitude))
    if not address:
        return False
    employee.default_address = address
    employee.save(update_fields=["default_address"])
    return True


class GeofenceValidationError(Exception):
    def __init__(self, message: str = "Attendance not marked.", admin_details: dict | None = None):
        super().__init__(message)
        self.message = message
        self.admin_details = admin_details or {}


def validate_geofence(
    employee: Employee,
    assignment: Assignment | None,
    latitude: float,
    longitude: float,
    accuracy: float | None = None
) -> None:
    if assignment:
        raw_radius = float(assignment.radius or settings.DEFAULT_GEOFENCE_RADIUS_METERS)
        radius = raw_radius * 1000 if raw_radius <= 10 else raw_radius
        target_lat = float(assignment.latitude)
        target_lon = float(assignment.longitude)
        location_label = assignment.patient_address or assignment.patient_name
    else:
        if employee.default_latitude is None or employee.default_longitude is None:
            if not employee.default_address or not ensure_default_coordinates(employee):
                logger.info("No active assignment or default work location for employee %s; skipping geofence validation.", employee.employee_id)
                return
        target_lat = float(employee.default_latitude)
        target_lon = float(employee.default_longitude)
        raw_radius = float(employee.default_radius or 0.1)
        radius = raw_radius * 1000 if raw_radius <= 10 else raw_radius
        location_label = employee.default_address or "Default Profile Location"

    effective_radius = max(500.0, radius)
    distance = distance_meters(float(latitude), float(longitude), target_lat, target_lon)
    buffer = max(100.0, (accuracy or 0) * 2.0)

    if distance > effective_radius + buffer:
        dist_diff_km = round(distance / 1000, 2)
        logger.warning(
            "\n================ [ADMIN GEOFENCE LOG] ================\n"
            "Employee:\n%s (%s)\n\n"
            "Assigned Location:\n%s\n\n"
            "Current Distance:\n%.2f km\n\n"
            "Reason:\nOutside geofence (Allowed radius: %.0f meters)\n\n"
            "Timestamp:\n%s\n"
            "======================================================",
            employee.name, employee.employee_id, location_label, dist_diff_km, effective_radius, timezone.now()
        )
        raise GeofenceValidationError(
            message=f"You are {dist_diff_km} km away from {location_label}.",
            admin_details={
                "employee": f"{employee.name} ({employee.employee_id})",
                "assigned_location": location_label,
                "current_distance_km": dist_diff_km,
                "allowed_radius_meters": effective_radius,
                "reason": "Outside geofence",
                "timestamp": str(timezone.now()),
            }
        )


def upload_selfie(
    image_file,
    folder: str,
    *,
    employee: Employee | None = None,
    timestamp: str | None = None,
    location: str | None = None,
) -> dict[str, str]:
    from apps.accounts.models import Employee

    if not employee and hasattr(image_file, "employee") and isinstance(image_file.employee, Employee):
        employee = image_file.employee

    employee_id = employee.employee_id if employee else "unknown"
    employee_name = employee.name if employee else "Unknown"

    attendance_type = "checkin" if folder == "attendance" else "checkout"
    result = upload_attendance_image(
        image_file,
        employee_id=employee_id,
        attendance_type=attendance_type,
        timestamp=timestamp,
        address=location,
        employee_name=employee_name,
    )
    return {"url": result["url"], "public_id": result["public_id"]}


def upload_profile_photo(image_file, *, employee_id: str, employee_name: str | None = None) -> dict[str, str]:
    result = upload_profile_image(image_file, employee_id=employee_id, employee_name=employee_name)
    return {"url": result["url"], "public_id": result["public_id"]}


def get_employee_presence_summary(employee: Employee, *, reference_time: datetime | None = None) -> dict[str, Any]:
    reference_time = reference_time or timezone.now()
    local_ref = timezone.localtime(reference_time)
    today = local_ref.date()
    
    sessions = Session.objects.filter(employee=employee).order_by('-login_time')[:10]
    session = None
    for s in sessions:
        if s.login_time:
            s_local_date = timezone.localtime(s.login_time).date()
            if s_local_date == today:
                session = s
                break

    if not session:
        return {
            'is_present': False,
            'status': 'Absent',
            'check_in_time': None,
            'check_out_time': None,
            'session_duration_seconds': 0,
            'session': None,
        }

    if session.is_active or not session.logout_time:
        duration_seconds = max(int((reference_time - session.login_time).total_seconds()), 0)
        return {
            'is_present': True,
            'status': 'Present',
            'check_in_time': session.login_time,
            'check_out_time': None,
            'session_duration_seconds': duration_seconds,
            'session': session,
        }

    logout_time = session.logout_time or session.login_time
    duration_seconds = max(int((logout_time - session.login_time).total_seconds()), 0)
    return {
        'is_present': True,
        'status': 'Checked Out',
        'check_in_time': session.login_time,
        'check_out_time': logout_time,
        'session_duration_seconds': duration_seconds,
        'session': session,
    }


def get_session_for_date(employee: Employee, target_date: date) -> Session | None:
    return (
        Session.objects.filter(employee=employee, login_time__date__lte=target_date)
        .filter(Q(logout_time__isnull=True) | Q(logout_time__date__gte=target_date))
        .order_by('-login_time')
        .first()
    )


def generate_attendance_export(start_date: date, end_date: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from datetime import time

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Report Header Block
    ws.cell(row=1, column=1, value="Skandan Home Carre Clinic LLP").font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    ws.cell(row=2, column=1, value="Attendance Report").font = Font(name="Calibri", size=13, bold=True, color="595959")
    ws.cell(row=3, column=1, value=f"Date Range: {start_date.strftime('%d-%b-%Y')} to {end_date.strftime('%d-%b-%Y')}").font = Font(name="Calibri", size=11, italic=True)
    ws.cell(row=4, column=1, value=f"Generated On: {timezone.localtime(timezone.now()).strftime('%d-%b-%Y %I:%M %p')} | Generated By: HR Admin").font = Font(name="Calibri", size=10, color="7F7F7F")

    # Generate Date list
    date_list = []
    curr = start_date
    while curr <= end_date:
        date_list.append(curr)
        curr += timedelta(days=1)

    # Styling definitions
    header_fill = PatternFill(start_color="6B2FA0", end_color="6B2FA0", fill_type="solid")
    sub_header_fill = PatternFill(start_color="F2EBF9", end_color="F2EBF9", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sub_header_font = Font(name="Calibri", size=10, bold=True, color="6B2FA0")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")

    # Row 6: Main Header Row
    ws.cell(row=6, column=1, value="Employee ID")
    ws.cell(row=6, column=2, value="Employee Name")
    ws.cell(row=6, column=3, value="Service Start")

    # Merge Row 6 & Row 7 for first three columns
    ws.merge_cells(start_row=6, start_column=1, end_row=7, end_column=1)
    ws.merge_cells(start_row=6, start_column=2, end_row=7, end_column=2)
    ws.merge_cells(start_row=6, start_column=3, end_row=7, end_column=3)

    col_idx = 4
    for dt in date_list:
        date_str = f"{dt.strftime('%d-%b-%Y')} ({dt.strftime('%a')})"
        ws.cell(row=6, column=col_idx, value=date_str)
        ws.merge_cells(start_row=6, start_column=col_idx, end_row=6, end_column=col_idx + 2)
        
        # Row 7: Sub-headers
        ws.cell(row=7, column=col_idx, value="Check In")
        ws.cell(row=7, column=col_idx + 1, value="Check Out")
        ws.cell(row=7, column=col_idx + 2, value="Working Hours")
        
        col_idx += 3

    # Last column: Working Days
    working_days_col = col_idx
    ws.cell(row=6, column=working_days_col, value="Working Days")
    ws.merge_cells(start_row=6, start_column=working_days_col, end_row=7, end_column=working_days_col)

    # Apply Header Styles
    for r in range(6, 8):
        for c in range(1, working_days_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = center_align
            cell.border = thin_border
            if r == 6:
                cell.fill = header_fill
                cell.font = header_font
            else:
                if c >= 4 and c < working_days_col:
                    cell.fill = sub_header_fill
                    cell.font = sub_header_font

    # Bulk query all sessions and attendances in date range (1 query each instead of 900+)
    tz = timezone.get_current_timezone()
    start_dt = timezone.make_aware(datetime.combine(start_date, time.min), tz)
    end_dt = timezone.make_aware(datetime.combine(end_date, time.max), tz)

    all_sessions = Session.objects.filter(login_time__gte=start_dt, login_time__lte=end_dt).order_by('login_time')
    sessions_by_emp_date: dict[tuple[int, date], Session] = {}
    for s in all_sessions:
        if s.login_time:
            loc_date = timezone.localtime(s.login_time).date()
            key = (s.employee_id, loc_date)
            if key not in sessions_by_emp_date or (s.is_active or s.logout_time):
                sessions_by_emp_date[key] = s

    all_attendances = Attendance.objects.filter(timestamp__gte=start_dt, timestamp__lte=end_dt).order_by('timestamp')
    attendances_by_emp_date: dict[tuple[int, date], Attendance] = {}
    for a in all_attendances:
        if a.timestamp:
            loc_date = timezone.localtime(a.timestamp).date()
            key = (a.employee_id, loc_date)
            if key not in attendances_by_emp_date:
                attendances_by_emp_date[key] = a

    # Populate Data Rows
    employees = Employee.objects.all().order_by("employee_id")
    row_idx = 8
    now = timezone.now()

    for emp in employees:
        service_start_str = emp.created_at.strftime("%d-%b-%Y") if emp.created_at else "-"
        ws.cell(row=row_idx, column=1, value=emp.employee_id).alignment = left_align
        ws.cell(row=row_idx, column=2, value=emp.name).alignment = left_align
        ws.cell(row=row_idx, column=3, value=service_start_str).alignment = center_align

        ws.cell(row=row_idx, column=1).font = bold_font
        ws.cell(row=row_idx, column=2).font = bold_font
        ws.cell(row=row_idx, column=3).font = regular_font

        c_idx = 4
        present_count = 0
        for dt in date_list:
            session = sessions_by_emp_date.get((emp.id, dt))
            attendance_rec = attendances_by_emp_date.get((emp.id, dt))

            if not session and not attendance_rec:
                check_in_val = "-"
                check_out_val = "-"
                hours_val = "Absent"
            else:
                present_count += 1
                login_time = session.login_time if session else attendance_rec.timestamp
                logout_time = session.logout_time if session else None

                check_in_val = timezone.localtime(login_time).strftime("%I:%M %p") if login_time else "-"

                if session and (session.is_active or not logout_time):
                    check_out_val = "Active"
                    diff = now - login_time
                    h = int(diff.total_seconds() // 3600)
                    m = int((diff.total_seconds() % 3600) // 60)
                    hours_val = f"{h}h {m}m"
                elif logout_time:
                    check_out_val = timezone.localtime(logout_time).strftime("%I:%M %p")
                    diff = logout_time - login_time
                    h = int(diff.total_seconds() // 3600)
                    m = int((diff.total_seconds() % 3600) // 60)
                    hours_val = f"{h}h {m}m"
                else:
                    check_out_val = "-"
                    hours_val = "Present"

            c1 = ws.cell(row=row_idx, column=c_idx, value=check_in_val)
            c2 = ws.cell(row=row_idx, column=c_idx + 1, value=check_out_val)
            c3 = ws.cell(row=row_idx, column=c_idx + 2, value=hours_val)

            for c in (c1, c2, c3):
                c.alignment = center_align
                c.font = regular_font
                c.border = thin_border
                if hours_val == "Absent":
                    c.font = Font(name="Calibri", size=10, color="9C0006")

            c_idx += 3

        # Write Working Days count in the last column
        c_work = ws.cell(row=row_idx, column=working_days_col, value=present_count)
        c_work.alignment = center_align
        c_work.font = bold_font
        c_work.border = thin_border

        row_idx += 1

    # Apply borders to employee fixed columns in data region
    for r in range(8, row_idx):
        ws.cell(row=r, column=1).border = thin_border
        ws.cell(row=r, column=2).border = thin_border
        ws.cell(row=r, column=3).border = thin_border
        ws.cell(row=r, column=working_days_col).border = thin_border

    # Freeze Panes below headers and after column C
    ws.freeze_panes = "D8"

    # Set column widths efficiently
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 16
    for c in range(4, working_days_col):
        col_letter = get_column_letter(c)
        ws.column_dimensions[col_letter].width = 14
    ws.column_dimensions[get_column_letter(working_days_col)].width = 16

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@transaction.atomic
def start_session(employee: Employee) -> Session:
    today = timezone.localdate()
    # Close any stale active sessions from past days
    Session.objects.select_for_update().filter(
        employee=employee, is_active=True, login_time__date__lt=today
    ).update(is_active=False)

    active_session = Session.objects.select_for_update().filter(
        employee=employee, is_active=True, login_time__date=today
    ).first()
    if active_session:
        return active_session
    return Session.objects.create(employee=employee, is_active=True)


@transaction.atomic
def end_session(employee: Employee) -> Session:
    session = Session.objects.select_for_update().filter(employee=employee, is_active=True).first()
    if not session:
        raise ValidationError({"detail": "No active session found."})
    session.is_active = False
    session.logout_time = timezone.now()
    session.save(update_fields=["is_active", "logout_time"])
    return session


def record_attendance(
    *,
    employee: Employee,
    assignment: Assignment | None,
    session: Session | None,
    attendance_type: str,
    photo_url: str,
    photo_public_id: str = "",
    latitude: float,
    longitude: float,
    address: str,
    status: str,
    remarks: str = "",
) -> Attendance:
    return Attendance.objects.create(
        employee=employee,
        assignment=assignment,
        session=session,
        attendance_type=attendance_type,
        photo_url=photo_url,
        photo_public_id=photo_public_id,
        latitude=latitude,
        longitude=longitude,
        address=address,
        status=status,
        remarks=remarks,
    )


def annotate_image(image_file, *, timestamp: str | None = None, location: str | None = None):
    from PIL import Image as PILImage, ImageDraw, ImageFont

    image = PILImage.open(image_file).convert("RGBA")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()

    text = []
    if timestamp:
        text.append(timestamp)
    if location:
        text.append(location)
    if text:
        draw.text((10, 10), "\n".join(text), fill=(255, 255, 255, 255), font=font)
    output = BytesIO()
    image.convert("RGB").save(output, format="JPEG")
    output.seek(0)
    return output


def validate_liveness(image_file, liveness_score: float | None = None) -> None:
    if not liveness_service.is_live(image_file=image_file, liveness_score=liveness_score):
        raise ValidationError({"detail": "Liveness check failed."})


def log_location(
    *,
    session: Session,
    employee: Employee,
    latitude: float,
    longitude: float,
    accuracy: float | None = None,
    speed: float | None = None,
    battery_percentage: int | None = None,
    is_mock: bool = False,
    client_timestamp: datetime | None = None,
) -> LocationLog:
    if is_mock:
        raise ValidationError({"detail": "Mock location detected."})
    if client_timestamp:
        skew = abs((timezone.now() - client_timestamp).total_seconds())
        if skew > 300:
            raise ValidationError({"detail": "Timestamp skew is too large."})
    return LocationLog.objects.create(
        session=session,
        employee=employee,
        latitude=latitude,
        longitude=longitude,
        accuracy=accuracy,
        speed=speed,
        battery_percentage=battery_percentage,
    )
